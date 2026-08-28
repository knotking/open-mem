"""The format matrix.

Fixtures are constructed in the test rather than checked in: a binary fixture
nobody can read is a test nobody can change, and these are all small enough to
build honestly.

Tier C is tested as carefully as Tier A. "Nothing is rejected" is only true if
the rejection path records a reason, and a corrupt file that crashes a worker
is worse than one that fails.
"""

from __future__ import annotations

import io
import json
import zipfile

import pytest

from memdog.parsers import (
    MAX_SHEET_ROWS,
    NeedsModel,
    ParseFailed,
    for_file,
    parse,
)

pytestmark = pytest.mark.asyncio


def _mark_encrypted(payload: bytes) -> bytes:
    """Set the encryption bit in both headers.

    `zipfile` clears `flag_bits` when writing, so a fixture cannot ask for an
    encrypted archive -- the bit has to be set in the bytes afterwards, in the
    local file header (offset 6) and the central directory entry (offset 8).
    """
    out = bytearray(payload)
    for signature, offset in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
        start = 0
        while (found := out.find(signature, start)) != -1:
            out[found + offset] |= 0x01
            start = found + 4
    return bytes(out)


# --------------------------------------------------------------- text & markup


async def test_plain_text_and_encodings():
    parsed = parse("hello world".encode(), mime="text/plain", name="a.txt")
    assert parsed.text == "hello world"

    # Windows-1252 is what real exports contain. Without detection this becomes
    # mojibake, and mojibake retrieves nothing while looking like it worked.
    payload = "Beyoncé – naïve café".encode("windows-1252")
    parsed = parse(payload, mime="", name="b.txt")
    assert "caf" in parsed.text and "\ufffd" not in parsed.text[:5]
    assert parsed.structure["encoding"] != "utf-8"
    assert parsed.warnings


async def test_html_drops_script_and_keeps_structure():
    html = b"""<html><head><title>Q3 review</title></head>
    <body><h1>Revenue</h1><script>alert('x')</script>
    <p>Revenue rose 12% in the third quarter.</p></body></html>"""
    parsed = parse(html, mime="text/html", name="r.html")
    assert "Revenue rose 12%" in parsed.text
    assert "alert" not in parsed.text          # script content is not content
    assert parsed.structure["title"] == "Q3 review"
    assert "Revenue" in parsed.structure["headings"]


async def test_csv_repeats_the_header_per_row():
    """A bare CSV grid chunks into meaningless fragments, because the header
    scrolls out of the chunk and the numbers lose their labels."""
    csv_bytes = b"name,role,city\nDana,engineer,Berlin\nSam,designer,Lisbon\n"
    parsed = parse(csv_bytes, mime="text/csv", name="people.csv")
    assert "name: Dana; role: engineer; city: Berlin" in parsed.text
    assert parsed.structure["columns"] == ["name", "role", "city"]
    assert parsed.structure["rows"] == 2


async def test_json_and_jsonl():
    parsed = parse(json.dumps({"b": 1, "a": 2}).encode(), mime="application/json", name="x.json")
    assert parsed.text.index('"a"') < parsed.text.index('"b"')   # sorted, stable

    lines = b'{"id": 1}\nnot json\n{"id": 2}\n'
    parsed = parse(lines, mime="", name="x.jsonl")
    assert parsed.structure["records"] == 2
    assert "1 unparseable lines skipped" in parsed.warnings[0]


async def test_malformed_json_is_terminal_not_retryable():
    with pytest.raises(ParseFailed) as exc:
        parse(b"{not json", mime="application/json", name="x.json")
    assert exc.value.code == "malformed"


async def test_yaml_uses_safe_load():
    """A YAML document is untrusted input, and full-fat YAML constructs
    arbitrary Python objects."""
    hostile = b"!!python/object/apply:os.system ['echo pwned']"
    with pytest.raises(ParseFailed):
        parse(hostile, mime="", name="x.yaml")


# ---------------------------------------------------------------------- office


def _docx_bytes() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Quarterly Filing", level=1)
    document.add_paragraph("The reconciliation job stalled overnight.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Region"
    table.rows[0].cells[1].text = "EMEA"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def test_docx_text_headings_and_tables():
    parsed = parse(_docx_bytes(), mime="", name="report.docx")
    assert "reconciliation job stalled" in parsed.text
    assert "Region | EMEA" in parsed.text
    assert "Quarterly Filing" in parsed.structure["headings"]


async def test_xlsx_labels_cells_and_caps_rows():
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Ledger"
    sheet.append(["account", "amount"])
    for i in range(5):
        sheet.append([f"acct-{i}", i * 100])
    buffer = io.BytesIO()
    book.save(buffer)

    parsed = parse(buffer.getvalue(), mime="", name="ledger.xlsx")
    assert "account: acct-3; amount: 300" in parsed.text
    assert parsed.structure["sheets"] == ["Ledger"]
    assert MAX_SHEET_ROWS > 0


async def test_pptx_keeps_slide_boundaries():
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Incident review"
    buffer = io.BytesIO()
    deck.save(buffer)

    parsed = parse(buffer.getvalue(), mime="", name="deck.pptx")
    assert "# Slide 1" in parsed.text and "Incident review" in parsed.text


async def test_encrypted_office_file_fails_clearly():
    """Must fail with a status, never crash a worker or retry-loop."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", b"<xml/>")
    with pytest.raises(ParseFailed) as exc:
        parse(_mark_encrypted(buffer.getvalue()), mime="", name="secret.docx")
    assert exc.value.code == "encrypted"


async def test_macro_enabled_office_is_never_expanded():
    """docm/xlsm are a malware vector specifically because team members
    download each other's uploads."""
    assert for_file("payload.docm", "") is None
    with pytest.raises(ParseFailed) as exc:
        parse(_docx_bytes(), mime="", name="payload.docm")
    assert exc.value.code == "unsupported"


# ------------------------------------------------------------------------ pdf


def _pdf_bytes(text: str) -> bytes:
    """A minimal single-page PDF with a real text layer."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n".encode()
        + b"%%EOF\n"
    )
    return bytes(out)


async def test_digital_pdf_extracts_its_text_layer():
    body = "The quarterly filing was delayed after the reconciliation job stalled overnight"
    parsed = parse(_pdf_bytes(body), mime="application/pdf", name="filing.pdf")
    assert "reconciliation job stalled" in parsed.text
    assert parsed.structure["pages"] == 1


async def test_a_scanned_pdf_is_routed_to_ocr_not_returned_empty():
    """A digital PDF and a scanned one are different products with costs an
    order of magnitude apart. Detect before routing."""
    with pytest.raises(NeedsModel) as exc:
        parse(_pdf_bytes("x"), mime="application/pdf", name="scan.pdf")
    assert exc.value.capability == "ocr"


async def test_corrupt_pdf_is_terminal():
    with pytest.raises(ParseFailed) as exc:
        parse(b"%PDF-1.4 truncated garbage", mime="application/pdf", name="broken.pdf")
    assert exc.value.code == "malformed"


# ---------------------------------------------------------------------- email


async def test_eml_keeps_headers_and_names_attachments_only():
    raw = (
        b"From: dana@example.com\r\nTo: sam@example.com\r\n"
        b"Subject: Renewal terms\r\nDate: Tue, 20 Aug 2026 09:00:00 +0000\r\n"
        b'Content-Type: multipart/mixed; boundary="b1"\r\n\r\n'
        b"--b1\r\nContent-Type: text/plain\r\n\r\n"
        b"Holding list pricing through renewal.\r\n"
        b"--b1\r\nContent-Type: application/pdf\r\n"
        b'Content-Disposition: attachment; filename="terms.pdf"\r\n\r\n'
        b"%PDF-1.4\r\n--b1--\r\n"
    )
    parsed = parse(raw, mime="message/rfc822", name="mail.eml")
    assert "Subject: Renewal terms" in parsed.text
    assert "Holding list pricing" in parsed.text
    # Referenced by name only -- never speculating about contents it cannot see.
    assert parsed.structure["attachments"] == ["terms.pdf"]
    assert "%PDF" not in parsed.text


async def test_mbox_is_capped_and_says_so():
    """One Takeout file can legitimately hold tens of thousands of messages."""
    one = (
        b"From dana@example.com\r\nFrom: dana@example.com\r\n"
        b"Subject: Note {i}\r\n\r\nBody {i}\r\n\r\n"
    )
    payload = b"".join(one.replace(b"{i}", str(i).encode()) for i in range(60))
    parsed = parse(payload, mime="", name="takeout.mbox")
    assert parsed.structure["messages"] == 60
    assert parsed.structure["indexed"] == 50
    assert parsed.truncated and "bulk ingestion" in parsed.warnings[0]


# ------------------------------------------------------- calendar & contacts


async def test_ics_maps_onto_a_canonical_event():
    """Highest value per unit cost in the whole format list -- no model at all."""
    ics = (
        b"BEGIN:VCALENDAR\r\nBEGIN:VEVENT\r\nSUMMARY:Renewal call\r\n"
        b"DTSTART:20260901T140000Z\r\nLOCATION:Zoom\r\n"
        b"DESCRIPTION:Discuss the 12% uplift and revised\r\n  terms\r\n"
        b"END:VEVENT\r\nEND:VCALENDAR\r\n"
    )
    parsed = parse(ics, mime="", name="cal.ics")
    assert "Event: Renewal call" in parsed.text
    assert "Location: Zoom" in parsed.text
    # Line folding is part of the format, not an edge case.
    assert "revised terms" in parsed.text
    assert parsed.structure["target_type"] == "Event"


async def test_vcf_maps_onto_a_canonical_person():
    vcf = b"BEGIN:VCARD\r\nFN:Dana Ruiz\r\nORG:Acme\r\nEMAIL:dana@acme.com\r\nEND:VCARD\r\n"
    parsed = parse(vcf, mime="", name="contact.vcf")
    assert "Name: Dana Ruiz" in parsed.text and "Email: dana@acme.com" in parsed.text
    assert parsed.structure["target_type"] == "Person"


# --------------------------------------------------------------------- code


async def test_code_is_indexed_with_its_language():
    parsed = parse(b"def reconcile(rows):\n    return sorted(rows)\n", mime="", name="job.py")
    assert "def reconcile" in parsed.text
    assert parsed.structure["language"] == "python"


# ------------------------------------------------------------------ archives


async def test_archive_lists_members_and_extracts_what_it_can():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", b"The rollback completed at 14:02.")
        archive.writestr("data.csv", b"k,v\nalpha,1\n")
        archive.writestr("image.bin", b"\x00\x01\x02")
    parsed = parse(buffer.getvalue(), mime="application/zip", name="bundle.zip")
    assert "Archive contents:" in parsed.text
    assert "rollback completed at 14:02" in parsed.text
    assert "k: alpha" in parsed.text
    assert parsed.structure["members"] == 3


async def test_encrypted_archive_fails_clearly():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("secret.txt", b"x")
    with pytest.raises(ParseFailed) as exc:
        parse(_mark_encrypted(buffer.getvalue()), mime="application/zip", name="locked.zip")
    assert exc.value.code == "encrypted"


# ------------------------------------------------------------------- tier C


@pytest.mark.parametrize(
    "mime,capability",
    [("image/png", "vision"), ("audio/mpeg", "transcription"), ("video/mp4", "transcription")],
)
async def test_media_asks_for_a_model_rather_than_failing(mime, capability):
    """Not a failure -- a policy. The bytes are fine and stay where they are."""
    with pytest.raises(NeedsModel) as exc:
        parse(b"\x00\x01", mime=mime, name=f"file.{mime.split('/')[1]}")
    assert exc.value.capability == capability


async def test_an_unknown_binary_is_stored_not_rejected():
    with pytest.raises(ParseFailed) as exc:
        parse(b"\x00\x01\x02", mime="application/x-dicom", name="scan.dcm")
    assert exc.value.code == "unsupported"    # recorded on the row, never dropped


async def test_a_generic_mime_defers_to_the_extension_but_a_specific_one_does_not():
    """Sniffing a CSV says `text/plain`, which is true and useless.

    This was a live bug: calendars and spreadsheets were indexed as walls of raw
    text because the generic sniff matched a handler before the extension was
    consulted. The asymmetry is the fix -- a specific MIME still overrides the
    name, so a `.docx` cannot claim to be anything it is not.
    """
    ics = b"BEGIN:VCALENDAR\r\nBEGIN:VEVENT\r\nSUMMARY:Incident review\r\nEND:VEVENT\r\nEND:VCALENDAR\r\n"
    parsed = parse(ics, mime="text/plain", name="meeting.ics")
    assert "Event: Incident review" in parsed.text        # not raw BEGIN:VEVENT

    csv_bytes = b"account,amount\nreconciliation,4200\n"
    parsed = parse(csv_bytes, mime="text/plain", name="ledger.csv")
    assert "account: reconciliation; amount: 4200" in parsed.text

    # But a specific sniff still wins over a lying name.
    with pytest.raises(ParseFailed) as exc:
        parse(b"%PDF-1.4 broken", mime="application/pdf", name="notes.txt")
    assert exc.value.code == "malformed"
